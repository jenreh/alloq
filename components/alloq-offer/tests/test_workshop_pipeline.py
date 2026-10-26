"""The one-click draft end to end: uploaded files to a downloadable Word draft."""

import io
from datetime import UTC, datetime, timedelta
from email.message import EmailMessage

import pytest
from alloq_offer.workshop.ai.agents import NOT_CONFIGURED
from alloq_offer.workshop.ai.demo import DemoOfferAi
from alloq_offer.workshop.ai.port import (
    AiSetupError,
    AiUnavailableError,
    BlockText,
    MapRequest,
    SourceText,
)
from alloq_offer.workshop.configuration import WorkshopConfig
from alloq_offer.workshop.content.text import stored_lines
from alloq_offer.workshop.entities import (
    ChapterEntity,
    ExportEntity,
    ExportStatus,
    SourceBlockEntity,
    SourceKind,
    SourceStatus,
    StatementEntity,
    StatementType,
    StoryEntity,
)
from alloq_offer.workshop.jobs import queue
from alloq_offer.workshop.jobs.ai_access import failure_message, job_ai
from alloq_offer.workshop.jobs.draft_chapters import DraftChapters
from alloq_offer.workshop.jobs.export_docx import ExportDocx
from alloq_offer.workshop.jobs.extract_source import SourceExtractor
from alloq_offer.workshop.jobs.initial_draft import InitialDraft
from alloq_offer.workshop.jobs.registry import PermanentJobError
from alloq_offer.workshop.jobs.review_offer import ReviewOffer
from alloq_offer.workshop.jobs.runner import _context
from alloq_offer.workshop.repositories import workspace_repo
from alloq_offer.workshop.services import draft_overview, exports
from alloq_offer.workshop.services.chapter_draft import DRAFT_CHAPTERS_JOB
from alloq_offer.workshop.services.chapters import ChapterService, revision_of
from alloq_offer.workshop.services.exports import EXPORT_JOB
from alloq_offer.workshop.services.prompts import NO_PROMPT
from alloq_offer.workshop.services.review_request import REVIEW_JOB
from alloq_offer.workshop.services.source_revocation import revoke_source
from alloq_offer.workshop.services.sources import (
    EXTRACT_JOB,
    INITIAL_DRAFT_JOB,
    add_source,
    request_initial_draft,
)
from alloq_offer.workshop.services.workspace import open_workspace
from alloq_offer.workshop.storage import DbBlobStore
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import select

NOW = datetime(2026, 10, 24, tzinfo=UTC)
DOCX_REQUIREMENT = "Kunden sollen ihre Zählerstände online melden können."
PDF_REQUIREMENT = "Die Sachbearbeitung muss Abschlagszahlungen pruefen koennen."
PPTX_REQUIREMENT = "Kunden wollen Rechnungen als PDF herunterladen."
EML_QUESTION = "Ist eine Anbindung an SAP IS-U vorgesehen?"


class _NoTranscriber:
    async def transcribe(self, data: bytes, filename: str, mime: str) -> list:  # noqa: ARG002
        raise AssertionError


def _pdf(text: str) -> bytes:
    writer = PdfWriter()
    font = writer._add_object(  # noqa: SLF001
        DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
    )
    page = writer.add_blank_page(400, 200)
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 10 Tf 10 100 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)  # noqa: SLF001
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _docx() -> bytes:
    document = Document()
    document.add_paragraph(DOCX_REQUIREMENT)
    document.add_paragraph("Heute melden Kunden ihre Stände per Postkarte.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pptx() -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Rechnungen"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box.text_frame.text = PPTX_REQUIREMENT
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def _eml() -> bytes:
    message = EmailMessage()
    message["Subject"] = "Rückfragen zum Kundenportal"
    message["From"] = "it@stadtwerke.example"
    message["Message-ID"] = "<1@stadtwerke.example>"
    message.set_content(f"Guten Tag,\n\n{EML_QUESTION}\n")
    return message.as_bytes()


FILES = (
    ("Anforderungen.docx", SourceKind.DOCX, _docx),
    ("Fachkonzept.pdf", SourceKind.PDF, lambda: _pdf(PDF_REQUIREMENT)),
    ("Workshop.pptx", SourceKind.PPTX, _pptx),
    ("Rueckfragen.eml", SourceKind.EML, _eml),
)


class Pipeline:
    """Runs queued jobs of the test database with the demo AI, like the worker."""

    def __init__(self, db) -> None:  # noqa: ANN001
        self.db = db
        self.store = DbBlobStore(10_000_000, db.session)
        ai = DemoOfferAi()
        self.stages = {
            EXTRACT_JOB: SourceExtractor(
                self.store, _NoTranscriber(), db.session, 3, lambda: NOW
            ),
            INITIAL_DRAFT_JOB: InitialDraft(ai, db.session),
            DRAFT_CHAPTERS_JOB: DraftChapters(ai, db.session),
            REVIEW_JOB: ReviewOffer(ai, db.session),
            EXPORT_JOB: ExportDocx(self.store, db.session),
        }

    async def drain(self) -> list[str]:
        ran: list[str] = []
        for _ in range(50):
            async with self.db.session() as session:
                # Follow-up stages are queued with the real clock, not NOW.
                job = await queue.claim_next(
                    session,
                    now=datetime.now(UTC) + timedelta(days=1),
                    lease_seconds=60,
                )
                ctx = _context(job) if job is not None else None
            if ctx is None:
                return ran
            result = await self.stages[ctx.kind].run(ctx)
            async with self.db.session() as session:
                await queue.complete(session, ctx.job_id, ctx.attempt, result)
            ran.append(ctx.kind)
        raise AssertionError(ran)

    async def upload(self, auto_draft: bool = False) -> None:
        async with self.db.session() as session:
            workspace = await open_workspace(session, 1, "Anna")
            workspace.auto_draft = auto_draft
        for filename, kind, build in FILES:
            await add_source(
                self.store,
                offer_id=1,
                kind=kind,
                title=filename,
                filename=filename,
                mime="application/octet-stream",
                data=build(),
                user="Anna",
                now=NOW,
                max_attempts=3,
                session_factory=self.db.session,
            )


async def _download_text(db, store: DbBlobStore) -> str:  # noqa: ANN001
    async with db.session() as session:
        export_id = await session.scalar(
            select(ExportEntity.id)
            .where(ExportEntity.status == ExportStatus.READY)
            .order_by(ExportEntity.id.desc())
        )
    assert export_id is not None
    filename, data = await exports.load_download(db.session, store, 1, export_id)
    assert filename.endswith(".docx")
    document = Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs)


async def test_one_click_turns_uploaded_files_into_a_downloadable_word_draft(
    db,  # noqa: ANN001
) -> None:
    pipeline = Pipeline(db)
    await pipeline.upload()
    assert await pipeline.drain() == [EXTRACT_JOB] * len(FILES)

    await request_initial_draft(db.session, 1, now=NOW, max_attempts=3)
    assert await pipeline.drain() == [
        INITIAL_DRAFT_JOB,
        DRAFT_CHAPTERS_JOB,
        REVIEW_JOB,
        EXPORT_JOB,
    ]

    async with db.session() as session:
        stories = (await session.scalars(select(StoryEntity))).all()
        blocks = {
            b.id: b.text for b in (await session.scalars(select(SourceBlockEntity)))
        }
        statements = (await session.scalars(select(StatementEntity))).all()
        chapters = {
            c.chapter_key: c for c in (await session.scalars(select(ChapterEntity)))
        }
        summary = await revision_of(session, chapters["leistungsumfang"])
    # Stories quote the uploaded documents and cite the block they came from.
    goals = {s.goal: s for s in stories}
    story = goals["ihre Zählerstände online melden können"]
    assert DOCX_REQUIREMENT in blocks[story.source_refs[0]]
    assert "Rechnungen als PDF herunterladen" in goals
    assert "Abschlagszahlungen pruefen koennen" in goals
    assert any(
        s.type == StatementType.OPEN_POINT and s.text == EML_QUESTION
        for s in statements
    )
    # Prose chapters were written in the same run and kept as revisions.
    assert chapters["leistungsumfang"].revision > 1
    assert any(DOCX_REQUIREMENT in line for line in stored_lines(summary.content))
    # The Word draft is stored and downloadable with the chapter text.
    text = await _download_text(db, pipeline.store)
    assert DOCX_REQUIREMENT in text
    # Every live story is part of the offer, so the first draft lists it.
    assert "Ihre Zählerstände online melden können" in text
    async with db.session() as session:
        overview = await draft_overview.load_overview(session, 1)
    assert [stage["status"] for stage in overview.stages] == [draft_overview.DONE] * 5
    assert overview.draft_export_id > 0


async def test_auto_draft_starts_the_pipeline_when_sources_are_usable(db) -> None:  # noqa: ANN001
    pipeline = Pipeline(db)
    await pipeline.upload(auto_draft=True)
    ran = await pipeline.drain()
    assert ran[: len(FILES)] == [EXTRACT_JOB] * len(FILES)
    assert ran.count(INITIAL_DRAFT_JOB) == 1
    assert ran[-1] == EXPORT_JOB
    async with db.session() as session:
        workspace = await workspace_repo.find_by_offer(session, 1)
        assert workspace is not None
        assert workspace.map_revision == 1


async def test_without_auto_draft_usable_sources_only_wait(db) -> None:  # noqa: ANN001
    pipeline = Pipeline(db)
    await pipeline.upload(auto_draft=False)
    await pipeline.drain()
    async with db.session() as session:
        overview = await draft_overview.load_overview(session, 1)
    assert [s["status"] for s in overview.stages] == [
        draft_overview.DONE,
        *[draft_overview.OPEN] * 4,
    ]
    assert not overview.working


async def test_extraction_shows_the_processing_status_while_it_runs(db) -> None:  # noqa: ANN001
    pipeline = Pipeline(db)
    await pipeline.upload()
    extractor = pipeline.stages[EXTRACT_JOB]
    assert isinstance(extractor, SourceExtractor)
    source = await extractor._begin(1, 1)  # noqa: SLF001
    assert source is not None
    async with db.session() as session:
        stored = await session.get(type(source), 1)
        assert stored is not None
        assert stored.status == SourceStatus.PROCESSING


async def test_revoking_a_source_keeps_an_export_with_the_same_bytes(db) -> None:  # noqa: ANN001
    store = DbBlobStore(1_000_000, db.session)
    data = _docx()
    added = await add_source(
        store,
        offer_id=1,
        kind=SourceKind.DOCX,
        title="a.docx",
        filename="a.docx",
        mime="application/octet-stream",
        data=data,
        user="Anna",
        now=NOW,
        max_attempts=3,
        session_factory=db.session,
    )
    async with db.session() as session:
        await open_workspace(session, 1, "Anna")
    await ChapterService(db.session, 1, "Anna").open()
    export = await exports.request_export(
        db.session,
        1,
        exports.ExportKind.DRAFT,
        user="Anna",
        confirmed=False,
        unsaved=False,
        now=NOW,
        max_attempts=3,
    )
    async with db.session() as session:
        row = await session.get(ExportEntity, export.id)
        assert row is not None
        row.blob_id = added.source.blob_id
    await revoke_source(db.session, 1, added.source.id, now=NOW)
    assert await store.get(added.source.blob_id) == data


def test_missing_ai_setup_fails_the_job_with_a_readable_message() -> None:
    with pytest.raises(PermanentJobError, match="nicht eingerichtet"):
        job_ai(WorkshopConfig())
    assert "Azure" in NOT_CONFIGURED
    assert isinstance(job_ai(WorkshopConfig(ai_provider="demo")), DemoOfferAi)
    setup = AiSetupError(NO_PROMPT)
    assert failure_message(setup, "Allgemein") == NO_PROMPT
    assert failure_message(AiUnavailableError("propose_map"), "Allgemein") == (
        "Allgemein"
    )


def _source(*texts: str, gaps: list[str] | None = None) -> SourceText:
    return SourceText(
        id=1,
        title="Notizen.txt",
        kind="txt",
        blocks=[BlockText(id=i, text=t) for i, t in enumerate(texts, 1)],
        gaps=gaps or [],
    )


async def test_demo_ai_quotes_sources_and_falls_back_without_requirements() -> None:
    ai = DemoOfferAi()
    proposed = await ai.propose_map(
        [_source("Ein Satz ohne Wunsch und ohne Frage.", gaps=["Folie 2: Bild"])]
    )
    story = proposed.activities[0].tasks[0].stories[0]
    assert story.goal == "die Inhalte aus „Notizen“ im System nutzen"
    assert story.skeleton
    assert proposed.open_questions[0].text == "In „Notizen“ fehlt: Folie 2: Bild"
    briefing = await ai.propose_briefing(
        [_source("Annahme: der Betrieb bleibt beim Kunden.", "Gibt es eine API?")]
    )
    assert [s.type for s in briefing.statements] == [
        StatementType.ASSUMPTION,
        StatementType.OPEN_POINT,
    ]


async def test_demo_ai_reads_lines_and_skips_mail_headers() -> None:
    briefing = await DemoOfferAi().propose_briefing(
        [
            _source(
                "Betreff: Kundenportal Anfrage\nFachkonzept Abrechnung\n"
                "Die Abrechnung muss Stände prüfen."
            )
        ]
    )
    assert [s.text for s in briefing.statements] == [
        "Fachkonzept Abrechnung",
        "Die Abrechnung muss Stände prüfen.",
    ]


async def test_demo_ai_map_additions_leave_out_what_the_map_has() -> None:
    ai = DemoOfferAi()
    source = _source(
        "Kunden sollen Termine buchen können.", "Kunden wollen Belege sehen."
    )
    request = MapRequest(
        instruction="",
        map_outline="Als Kunden möchte ich Termine buchen können, um ...",
        statements=[],
        sources=[source],
    )
    proposed = await ai.propose_map_additions(request)
    goals = [s.goal for a in proposed.activities for t in a.tasks for s in t.stories]
    assert goals == ["Belege sehen"]
