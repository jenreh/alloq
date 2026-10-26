"""Offer administration: block, template and prompt states, components and page."""

import io
import json
from collections.abc import Iterator
from contextlib import ExitStack, contextmanager
from typing import Any
from unittest.mock import AsyncMock, MagicMock, patch

import reflex as rx
from alloq_offer.admin import pages
from alloq_offer.admin.components import blocks, prompts, templates
from alloq_offer.admin.states import (
    block_admin_state,
    prompt_admin_state,
    template_admin_state,
)
from alloq_offer.admin.states.block_admin_state import BlockAdminState
from alloq_offer.admin.states.offer_admin_state import TAB_BLOCKS, OfferAdminState
from alloq_offer.admin.states.prompt_admin_state import PromptAdminState
from alloq_offer.admin.states.template_admin_state import TemplateAdminState
from alloq_offer.workshop.configuration import WorkshopConfig
from alloq_offer.workshop.export.template_package import CHAPTER_ANCHOR
from alloq_offer.workshop.services.template_drafts import NOT_DOCX
from docx import Document

from appkit_user.authentication.decorators import _NOT_AUTHORIZED_MSG
from appkit_user.authentication.states import LoginState

MODULES = (block_admin_state, template_admin_state, prompt_admin_state)


def _u(text: str) -> str:
    return json.dumps(text)[1:-1]


def _shows(text: str, event: str) -> bool:
    return text in event or _u(text) in event


@contextmanager
def _as_admin(state: rx.State, db, *, is_admin: bool = True) -> Iterator[None]:  # noqa: ANN001
    user = MagicMock(is_admin=is_admin, user_id=1, roles=[])
    user.name = "Admin"

    async def _user() -> MagicMock:
        return user

    login = MagicMock()
    type(login).authenticated_user = property(lambda _self: _user())
    login.redir = AsyncMock(return_value=None)

    async def _get_state(cls: type) -> Any:
        assert cls is LoginState
        return login

    registry = MagicMock()
    registry.get.return_value = WorkshopConfig()
    object.__setattr__(state, "get_state", _get_state)
    with ExitStack() as stack:
        for module in MODULES:
            stack.enter_context(
                patch(f"{module.__name__}.get_asyncdb_session", db.session)
            )
        stack.enter_context(
            patch(
                f"{template_admin_state.__name__}.service_registry",
                return_value=registry,
            )
        )
        try:
            yield
        finally:
            object.__delattr__(state, "get_state")


async def _run(handler: Any, *args: Any) -> str:
    result = handler(*args)
    if hasattr(result, "__aiter__"):
        return " ".join([str(i) async for i in result if i is not None])
    result = await result
    return "" if result is None else str(result)


def _new(cls: type[rx.State]) -> Any:
    return cls(_reflex_internal_init=True)  # ty: ignore[unknown-argument]


def _layout() -> bytes:
    document = Document()
    document.add_heading("Hinweis", 1)
    document.add_paragraph(CHAPTER_ANCHOR)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


class _Upload:
    def __init__(self, name: str, data: bytes) -> None:
        self.name = name
        self.size = len(data)
        self._data = data

    async def read(self) -> bytes:
        return self._data


async def test_non_admins_are_refused(db) -> None:  # noqa: ANN001
    for cls, handler, args in (
        (BlockAdminState, "save", ()),
        (TemplateAdminState, "publish", ()),
        (PromptAdminState, "activate", (1,)),
    ):
        state = _new(cls)
        with _as_admin(state, db, is_admin=False):
            event = await _run(getattr(state, handler), *args)
        assert _shows(_NOT_AUTHORIZED_MSG, event), handler


async def test_block_is_created_from_the_editor(db) -> None:  # noqa: ANN001
    state = _new(BlockAdminState)
    with _as_admin(state, db):
        await _run(state.load_blocks)
        state.open_create()
        assert state.editor_open
        assert state.creating
        for name, value in (
            ("key", "glossar"),
            ("title", "Glossar"),
            ("chapter_type", "table"),
            ("level", "2"),
            ("columns", ["Begriff", "Bedeutung"]),
            ("content_rules", ["Kurz."]),
            ("checks", ["no_acceptance_criteria"]),
            ("unknown", "ignored"),
        ):
            state.set_field(name, value)
        toast = await _run(state.save)
    assert _shows(block_admin_state.SAVED, toast)
    assert not state.editor_open
    created = next(b for b in state.blocks if b["key"] == "glossar")
    assert created["default_content"]["columns"] == ["Begriff", "Bedeutung"]
    assert created["level"] == 2
    with _as_admin(state, db):
        state.open_edit("glossar")
        assert (state.columns, state.creating) == (["Begriff", "Bedeutung"], False)
        state.set_field("chapter_type", "calc")
        state.set_field("fields", ["offer.person_days"])
        await _run(state.save)
        changed = next(b for b in state.blocks if b["key"] == "glossar")
        assert changed["default_content"]["rows"][0]["label"] == "Personentage"
        toast = await _run(state.delete, "glossar")
    assert _shows(block_admin_state.DELETED, toast)
    assert "glossar" not in {b["key"] for b in state.blocks}


async def test_block_editor_shows_german_errors(db) -> None:  # noqa: ANN001
    state = _new(BlockAdminState)
    with _as_admin(state, db):
        await _run(state.load_blocks)
        state.open_create()
        state.set_field("key", "X Y")
        toast = await _run(state.save)
        assert "Schlüssel" in toast or _u("Schlüssel") in toast
        state.set_field("key", "notiz")
        state.set_field("title", "Notiz")
        state.set_field("content_html", "<h1>Zu groß</h1>")
        toast = await _run(state.save)
        assert "Ebene" in toast
        state.set_field("chapter_type", "table")
        state.set_field("columns", "keine Liste")
        assert state.columns == ["Punkt", "Beschreibung"]
        state.set_field("columns", [str(i) for i in range(9)])
        toast = await _run(state.save)
        assert _shows(block_admin_state.BAD_COLUMNS, toast)
        state.open_edit("laufzeit")
        state.set_field("key", "management_summary")
        assert state.key == "laufzeit"
        assert state.used_in == ["Angebot mit Story Map"]
        assert "Laufzeit" in state.content_html or state.content_html == ""
        toast = await _run(state.delete, "laufzeit")
    assert _shows("Deaktivieren", toast)
    state.close_editor()
    assert not state.editor_open


async def test_template_outline_upload_and_publish(db) -> None:  # noqa: ANN001
    state = _new(TemplateAdminState)
    with _as_admin(state, db):
        await _run(state.load_templates)
        assert state.template_id
        assert state.template["key"] == "quo_v1"
        assert state.library == []
        first = state.outline[0]["id"]
        await _run(state.set_entry_title, first, "Kurzfassung")
        await _run(state.set_entry_level, first, "2")
        await _run(state.set_required, first, False)
        entry = state.outline[0]
        assert (entry["title"], entry["level_override"], entry["required"]) == (
            "Kurzfassung",
            2,
            False,
        )
        await _run(state.set_entry_level, first, "")
        assert state.outline[0]["level_override"] == 0
        assert await _run(state.set_entry_title, first, "Kurzfassung") == ""
        assert await _run(state.set_entry_level, first, "x") == ""
        assert await _run(state.set_required, 999, True) == ""
        await _run(state.reorder, str(first), 2)
        assert state.outline[2]["id"] == first
        assert state.order_version == 1
        await _run(state.remove_entry, first)
        assert first not in state.entry_ids
        assert state.library[0]["value"] == "management_summary"
        await _run(state.add_block, "management_summary")
        toast = await _run(state.handle_upload, [_Upload("a.pdf", b"x")])
        assert _shows(NOT_DOCX, toast)
        toast = await _run(state.handle_upload, [_Upload("neu.docx", _layout())])
        assert _shows(template_admin_state.UPLOADED, toast)
        assert state.template["frame_before"] == ["Hinweis"]
        toast = await _run(state.publish)
        assert "Version 2" in toast
        assert state.versions[0]["number"] == 2
        await _run(state.select_template, "x")
        await _run(state.select_template, str(state.template_id))
        assert await _run(state.handle_upload, []) != ""


async def test_prompt_versions_are_saved_and_activated(db) -> None:  # noqa: ANN001
    state = _new(PromptAdminState)
    with _as_admin(state, db):
        await _run(state.load_prompts)
        assert state.selected["name"] == "chapter"
        assert not state.changed
        state.set_text(state.text.rstrip() + "\nNeu.")
        assert state.changed
        toast = await _run(state.save)
        assert _shows(prompt_admin_state.SAVED, toast)
        assert state.versions[0]["version"] == 2
        toast = await _run(state.activate, 1)
        assert _shows(prompt_admin_state.ACTIVATED, toast)
        assert state.selected["version"] == 1
        await _run(state.select_prompt, "review")
        assert state.versions[0]["version"] == 1
        await _run(state.select_prompt, "gibt_es_nicht")
        assert state.name == "review"
        state.set_text(" ")
        toast = await _run(state.save)
        assert "20.000" in toast
        assert await _run(state.activate, "x") == ""


def test_tab_state_and_page_render() -> None:
    state = _new(OfferAdminState)
    state.set_tab(TAB_BLOCKS)
    assert state.tab == TAB_BLOCKS
    page = str(pages.create_offer_admin_page(rx.box())())
    assert _shows(templates.PUBLISH, page)
    for component, texts in (
        (templates.templates_tab(), (templates.PUBLISH, templates.LAYOUT_TITLE)),
        (blocks.blocks_table(), (blocks.NEW_BLOCK, blocks.PROMPT_LABEL)),
        (prompts.prompts_tab(), (prompts.SAVE, prompts.ACTIVATE)),
    ):
        rendered = str(component)
        for text in texts:
            assert _shows(text, rendered), text
