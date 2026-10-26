"""Admin services: chapter blocks, template drafts and publishing, prompts."""

import io
import re

import pytest
from alloq_offer.workshop.export.template_package import CHAPTER_ANCHOR
from alloq_offer.workshop.services import chapter_blocks, prompts, template_drafts
from alloq_offer.workshop.services.chapter_blocks import BlockError, parse_block
from alloq_offer.workshop.services.chapters import ChapterService
from alloq_offer.workshop.services.prompts import PromptError
from alloq_offer.workshop.services.template_drafts import DraftError, TemplateDraft
from alloq_offer.workshop.services.templates import latest_template
from alloq_offer.workshop.services.workspace import open_workspace
from alloq_offer.workshop.storage import BlobNotFoundError, DbBlobStore
from alloq_offer.workshop.template_entities import OfferTemplateEntity
from docx import Document
from sqlalchemy import select

BLOCK = {
    "key": "datenschutz",
    "title": " Datenschutz ",
    "chapter_type": "prose",
    "level": 2,
    "default_content": {"chapter_type": "prose"},
    "guiding_question": "Wie schützen wir Daten?",
    "content_rules": ["Keine Zusagen.", "  "],
    "prompt_text": "Nenne die DSGVO.",
    "required_info": ["Speicherort"],
    "checks": ["no_acceptance_criteria"],
}


def _layout(*before: str) -> bytes:
    document = Document()
    for heading in before:
        document.add_heading(heading, 1)
    document.add_paragraph(CHAPTER_ANCHOR)
    document.add_heading("Unterschriften", 1)
    document.core_properties.author = "Erika Muster"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


async def _template_id(db) -> int:  # noqa: ANN001
    async with db.session() as session:
        template = await session.scalar(select(OfferTemplateEntity))
        assert template is not None
        return template.id


# --- blocks --------------------------------------------------------------------


def test_block_input_is_validated_in_german() -> None:
    block = parse_block(BLOCK)
    assert (block.title, block.content_rules) == ("Datenschutz", ["Keine Zusagen."])
    for change, label in (
        ({"key": "Daten Schutz"}, "Schlüssel"),
        ({"title": ""}, "Titel"),
        ({"checks": ["gibt_es_nicht"]}, "Prüfungen"),
        ({"default_content": {"chapter_type": "table", "columns": ["A"]}}, "Start"),
        ({"content_rules": ["x"] * 21}, "Regeln"),
    ):
        with pytest.raises(BlockError, match=label):
            parse_block({**BLOCK, **change})


async def test_blocks_are_created_changed_listed_and_deleted(db) -> None:  # noqa: ANN001
    await chapter_blocks.save_block(db.session, parse_block(BLOCK), create=True)
    with pytest.raises(BlockError, match=chapter_blocks.DUPLICATE):
        await chapter_blocks.save_block(db.session, parse_block(BLOCK), create=True)
    changed = parse_block(
        {**BLOCK, "title": "Datenschutz und Betrieb", "active": False}
    )
    await chapter_blocks.save_block(db.session, changed, create=False)
    rows = {b["key"]: b for b in await chapter_blocks.list_blocks(db.session)}
    assert rows["datenschutz"]["title"] == "Datenschutz und Betrieb"
    assert not rows["datenschutz"]["active"]
    assert rows["datenschutz"]["used_in"] == []
    assert rows["laufzeit"]["used_in"] == ["Angebot mit Story Map"]
    with pytest.raises(BlockError, match=chapter_blocks.IN_USE):
        await chapter_blocks.delete_block(db.session, "laufzeit")
    await chapter_blocks.delete_block(db.session, "datenschutz")
    with pytest.raises(BlockError, match=chapter_blocks.NOT_FOUND):
        await chapter_blocks.delete_block(db.session, "datenschutz")
    with pytest.raises(BlockError, match=chapter_blocks.NOT_FOUND):
        await chapter_blocks.save_block(db.session, parse_block(BLOCK), create=False)


# --- template drafts -----------------------------------------------------------


async def test_outline_entries_are_added_moved_changed_and_removed(db) -> None:  # noqa: ANN001
    template_id = await _template_id(db)
    draft = TemplateDraft(db.session, template_id)
    await chapter_blocks.save_block(db.session, parse_block(BLOCK), create=True)
    await draft.add_block("datenschutz")
    with pytest.raises(DraftError, match=template_drafts.BLOCK_IN_OUTLINE):
        await draft.add_block("datenschutz")
    with pytest.raises(DraftError, match=template_drafts.NO_BLOCK):
        await draft.add_block("gibt_es_nicht")
    outline = await template_drafts.load_outline(db.session, template_id)
    added = outline[-1]
    assert (added["block_key"], added["level"], added["required"]) == (
        "datenschutz",
        2,
        False,
    )
    await draft.move(added["id"], 0)
    await draft.update(added["id"], required=True, title=" Schutz ", level=1)
    outline = await template_drafts.load_outline(db.session, template_id)
    assert (outline[0]["title"], outline[0]["level"], outline[0]["required"]) == (
        "Schutz",
        1,
        True,
    )
    await draft.update(added["id"], title="Neu")
    await draft.update(added["id"], level=0)
    outline = await template_drafts.load_outline(db.session, template_id)
    assert (outline[0]["title"], outline[0]["level"], outline[0]["required"]) == (
        "Neu",
        2,
        True,
    )
    with pytest.raises(DraftError, match=template_drafts.BAD_LEVEL):
        await draft.update(added["id"], required=True, title="", level=3)
    with pytest.raises(DraftError, match=template_drafts.TITLE_TOO_LONG):
        await draft.update(added["id"], required=True, title="x" * 256, level=None)
    await draft.remove(added["id"])
    outline = await template_drafts.load_outline(db.session, template_id)
    assert outline[0]["block_key"] == "management_summary"
    with pytest.raises(DraftError, match=template_drafts.NO_ENTRY):
        await draft.remove(added["id"])
    with pytest.raises(DraftError, match=template_drafts.NOT_FOUND):
        await TemplateDraft(db.session, 999).add_block("laufzeit")


async def test_layout_upload_is_checked_cleaned_and_reads_the_frame(
    db,  # noqa: ANN001
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    template_id = await _template_id(db)
    draft = TemplateDraft(db.session, template_id)
    store = DbBlobStore(10_000_000, db.session)
    with pytest.raises(DraftError, match="Word"):
        await draft.upload_layout(store, "layout.pdf", b"%PDF")
    with pytest.raises(DraftError, match="Word"):
        await draft.upload_layout(store, "layout.docx", b"kein zip")
    with pytest.raises(DraftError, match="Anker"):
        document = Document()
        out = io.BytesIO()
        document.save(out)
        await draft.upload_layout(store, "layout.docx", out.getvalue())
    with pytest.raises(DraftError, match=template_drafts.TOO_LARGE):
        await draft.upload_layout(DbBlobStore(10, db.session), "l.docx", _layout())
    monkeypatch.setattr(template_drafts, "MAX_XML_PART_BYTES", 100)
    with pytest.raises(DraftError, match=template_drafts.TOO_LARGE):
        await draft.upload_layout(store, "l.docx", _layout())
    monkeypatch.undo()
    await draft.upload_layout(store, "l.docx", _layout())
    first = (await template_drafts.list_templates(db.session))[0]
    await draft.upload_layout(store, "Layout.DOCX", _layout("Hinweis"))
    # The replaced draft layout is gone; no published version used it.
    with pytest.raises(BlobNotFoundError):
        await store.get(first["draft_layout_blob_id"])
    templates = await template_drafts.list_templates(db.session)
    assert (templates[0]["frame_before"], templates[0]["frame_after"]) == (
        ["Hinweis"],
        ["Unterschriften"],
    )
    assert templates[0]["latest"] == 1


async def test_publish_makes_the_next_version_for_new_offers_only(db) -> None:  # noqa: ANN001
    async with db.session() as session:
        await open_workspace(session, 1, "Anna")
    template_id = await _template_id(db)
    draft = TemplateDraft(db.session, template_id)
    await chapter_blocks.save_block(db.session, parse_block(BLOCK), create=True)
    await draft.add_block("datenschutz")
    await draft.upload_layout(
        DbBlobStore(10_000_000, db.session), "layout.docx", _layout("Hinweis")
    )
    assert await draft.publish("Admin") == 2
    with pytest.raises(DraftError, match="Version 2"):
        await draft.publish("Admin")
    versions = await template_drafts.list_versions(db.session, template_id)
    assert (versions[0]["number"], versions[0]["published_by"]) == (2, "Admin")
    async with db.session() as session:
        latest = await latest_template(session, "quo_v1")
    assert latest.version == "2"
    assert latest.chapters[-1].key == "datenschutz"
    assert latest.chapters[-1].prompt_text == "Nenne die DSGVO."
    assert latest.frame_before == ["Hinweis"]
    # The offer made before keeps version 1 and its own outline.
    keys = [c["key"] for c in await ChapterService(db.session, 1, "Anna").open()]
    assert "datenschutz" not in keys


async def test_publish_needs_a_layout_and_chapters(db) -> None:  # noqa: ANN001
    async with db.session() as session:
        template = OfferTemplateEntity(key="leer", name="Leer")
        session.add(template)
        await session.flush()
        template_id = template.id
    draft = TemplateDraft(db.session, template_id)
    with pytest.raises(DraftError, match=template_drafts.NO_LAYOUT):
        await draft.publish("Admin")
    await draft.upload_layout(DbBlobStore(10_000_000, db.session), "l.docx", _layout())
    with pytest.raises(DraftError, match=template_drafts.EMPTY_OUTLINE):
        await draft.publish("Admin")


# --- prompts -------------------------------------------------------------------


async def test_prompts_get_versions_and_can_go_back(db) -> None:  # noqa: ANN001
    listed = {p["name"]: p for p in await prompts.list_prompts(db.session)}
    assert listed["chapter"]["version"] == 1
    assert listed["chapter"]["label"] == prompts.PROMPT_LABELS["chapter"]
    old = listed["chapter"]["text"]
    assert (
        await prompts.save_prompt(
            db.session, "chapter", old.rstrip() + "\nNeu.", "Admin"
        )
        == 2
    )
    assert (await prompts.load_prompt("chapter", db.session)).label == "chapter@2"
    versions = await prompts.prompt_versions(db.session, "chapter")
    assert [(v["version"], v["active"]) for v in versions] == [(2, True), (1, False)]
    assert versions[0]["diff"] == [{"kind": "added", "text": "Neu."}]
    assert (versions[1]["first"], versions[1]["diff"]) == (True, [])
    await prompts.activate_prompt(db.session, "chapter", 1)
    assert (await prompts.load_prompt("chapter", db.session)).label == "chapter@1"
    with pytest.raises(PromptError, match=prompts.UNKNOWN_VERSION):
        await prompts.activate_prompt(db.session, "chapter", 9)
    with pytest.raises(PromptError, match=prompts.UNKNOWN_PROMPT):
        await prompts.save_prompt(db.session, "gibt_es_nicht", "x", "Admin")
    with pytest.raises(PromptError, match=re.escape(prompts.EMPTY_TEXT)):
        await prompts.save_prompt(db.session, "chapter", "   ", "Admin")
