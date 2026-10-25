import io
from pathlib import Path

import pytest
from alloq_offer.workshop.content.schema import (
    BoundValue,
    BulletList,
    CalcContent,
    CalcRow,
    Heading,
    OrderedList,
    Paragraph,
    ProseContent,
    Run,
    StoriesContent,
    TableBlock,
    TableContent,
)
from alloq_offer.workshop.export.chapter_renderers import NO_STORIES_TEXT
from alloq_offer.workshop.export.docx_renderer import DRAFT_MARKER, render_document
from alloq_offer.workshop.export.models import (
    MISSING_MARK,
    UNCONFIRMED_MARK,
    ChapterInput,
    RenderContext,
    RenderError,
    ResolvedStory,
    ResolvedValue,
)
from alloq_offer.workshop.export.template_package import load_template
from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from jinja2 import UndefinedError

SPECIAL = "Größe <b> & „Zitat“ 5 € \x0b"
SAMPLE_PATH = (
    Path(__file__).resolve().parents[3]
    / "docs"
    / "features"
    / "offer-workshop-sample.docx"
)


def _runs(text: str, bold: bool = False) -> list[Run]:
    return [Run(text=text, bold=bold)]


def _chapters() -> list[ChapterInput]:
    wide = TableBlock(
        header=[_runs(f"Spalte {i}") for i in range(8)],
        rows=[[_runs(f"{r}-{c}") for c in range(8)] for r in range(3)],
    )
    prose = ProseContent(
        blocks=[
            Paragraph(runs=_runs("Langer Text " * 200) + _runs(SPECIAL, bold=True)),
            Heading(level=3, runs=_runs("Unterpunkt")),
            BulletList(items=[_runs("Punkt A"), _runs("Punkt B")]),
            OrderedList(items=[_runs("Schritt 1"), _runs("Schritt 2")]),
            wide,
            BoundValue(field="offer.recommended_price"),
        ]
    )
    return [
        ChapterInput("management_summary", "Management Summary", 1, prose),
        ChapterInput(
            "user_stories",
            "Angebotene User Stories",
            2,
            StoriesContent(group_by="activity"),
        ),
        ChapterInput(
            "mitwirkung",
            "Mitwirkungspflichten",
            2,
            TableContent(
                columns=["Mitwirkung", "Wer"], rows=[[_runs("Zugang"), _runs("Kunde")]]
            ),
        ),
        ChapterInput(
            "verguetung",
            "Vergütung",
            1,
            CalcContent(
                intro=_runs("Kalkulation:"),
                rows=[
                    CalcRow(label="Preis", field="offer.recommended_price"),
                    CalcRow(label="Tage", field="offer.person_days"),
                    CalcRow(label="Kosten", field="offer.production_cost"),
                ],
            ),
        ),
    ]


def _context(draft: bool = True) -> RenderContext:
    stories = [
        ResolvedStory(
            key=f"US-{i:02d}",
            title=f"Story {i}",
            sentence=f"Als Rolle möchte ich Ziel {i}, um Nutzen zu haben.",
            group=f"Aktivität {i // 10 + 1}",
        )
        for i in range(1, 41)
    ]
    return RenderContext(
        cover={
            "project_title": "Projekt <Ä>",
            "customer_name": "Kunde GmbH",
            "customer_street": "Straße 1",
            "customer_city": "12345 Stadt",
            "customer_short": "Kunde",
            "offer_date": "1. Januar 2027",
            "quote_id": "QUO-1",
            "validity_date": "1. März 2027",
        },
        stories=stories,
        values={
            "offer.recommended_price": ResolvedValue("10.000 €", confirmed=True),
            "offer.person_days": ResolvedValue("12", confirmed=False),
        },
        draft=draft,
    )


@pytest.fixture(scope="module")
def rendered() -> bytes:
    return render_document(load_template(), _chapters(), _context())


def _style_name(item: object) -> str:
    style = getattr(item, "style", None)
    return style.name if style is not None else ""


def _doc(data: bytes) -> DocxDocument:
    return Document(io.BytesIO(data))


def test_cover_and_tail_are_filled(rendered: bytes) -> None:
    text = "\n".join(p.text for p in _doc(rendered).paragraphs)
    assert "Projekt <Ä>" in text
    assert "Kunde GmbH" in text
    assert "Leipzig, 1. Januar 2027" in text
    assert "bis einschließlich zum 1. März 2027" in text
    assert "{{" not in text
    assert "[[KAPITEL]]" not in text


def test_chapter_order_and_headings(rendered: bytes) -> None:
    doc = _doc(rendered)
    headings = [
        (_style_name(p), p.text)
        for p in doc.paragraphs
        if _style_name(p).startswith("Heading")
    ]
    titles = [t for _, t in headings]
    order = [
        "Management Summary",
        "Angebotene User Stories",
        "Mitwirkungspflichten",
        "Vergütung",
        "Rechtliche Bestimmungen",
    ]
    positions = [titles.index(t) for t in order]
    assert positions == sorted(positions)
    assert ("Heading 2", "Angebotene User Stories") in headings
    assert ("Heading 3", "Unterpunkt") in headings


def test_prose_blocks_and_special_characters(rendered: bytes) -> None:
    doc = _doc(rendered)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Größe <b> & „Zitat“ 5 €" in text
    styles = {_style_name(p) for p in doc.paragraphs}
    assert {"Aufzählung 1", "List Number"} <= styles
    wide = next(t for t in doc.tables if t.rows[0].cells[0].text == "Spalte 0")
    assert len(wide.columns) == 8
    assert _style_name(wide) == "arvato-Tabelle blau"
    assert wide.rows[3].cells[7].text == "2-7"


def test_stories_chapter_written_stories_grouped(rendered: bytes) -> None:
    doc = _doc(rendered)
    texts = [p.text for p in doc.paragraphs]
    assert "US-01 – Story 1" in texts
    assert "US-05 – Story 5" in texts  # every live story belongs to the offer
    assert "Als Rolle möchte ich Ziel 1, um Nutzen zu haben." in texts
    assert "Aktivität 1" in texts
    assert "Aktivität 5" in texts


def test_calc_marks_unconfirmed_and_missing(rendered: bytes) -> None:
    doc = _doc(rendered)
    calc = next(t for t in doc.tables if t.rows[0].cells[0].text == "Position")
    cells = {r.cells[0].text: r.cells[1].text for r in calc.rows[1:]}
    assert cells["Preis"] == "10.000 €"
    assert cells["Tage"] == f"12 {UNCONFIRMED_MARK}"
    assert cells["Kosten"] == MISSING_MARK


def test_draft_marker_toc_and_no_map(rendered: bytes) -> None:
    doc = _doc(rendered)
    assert any(p.text == DRAFT_MARKER for p in doc.paragraphs)
    assert any(
        i.text and "TOC" in i.text for i in doc.element.body.iter(qn("w:instrText"))
    )
    assert not any("Story Map" in p.text for p in doc.paragraphs)
    final = render_document(load_template(), _chapters(), _context(draft=False))
    assert not any(p.text == DRAFT_MARKER for p in _doc(final).paragraphs)


def test_a_map_without_stories_is_visible() -> None:
    chapters = [ChapterInput("s", "Stories", 2, StoriesContent())]
    ctx = RenderContext(cover=_context().cover, stories=[])
    doc = _doc(render_document(load_template(), chapters, ctx))
    assert any(p.text == NO_STORIES_TEXT for p in doc.paragraphs)


def test_unsupported_content_raises() -> None:
    bad = ChapterInput("x", "X", 1, ProseContent.model_construct(chapter_type="map"))  # type: ignore[arg-type]
    with pytest.raises(RenderError):
        render_document(load_template(), [bad], _context())


def test_rendered_file_reopens(rendered: bytes, tmp_path: Path) -> None:
    target = tmp_path / "sample.docx"
    target.write_bytes(rendered)
    assert _doc(target.read_bytes()).paragraphs


def test_stories_of_one_group_stay_together_when_unsorted() -> None:
    stories = [
        ResolvedStory("US-01", "A", "Als A möchte ich a, um x.", "G1"),
        ResolvedStory("US-02", "B", "Als B möchte ich b, um y.", "G2"),
        ResolvedStory("US-03", "C", "Als C möchte ich c, um z.", "G1"),
    ]
    chapters = [ChapterInput("s", "Stories", 2, StoriesContent())]
    ctx = RenderContext(cover=_context().cover, stories=stories)
    doc = _doc(render_document(load_template(), chapters, ctx))
    headings = [p.text for p in doc.paragraphs if p.text in {"G1", "G2"}]
    assert headings == ["G1", "G2"]


def test_missing_cover_value_fails_loudly() -> None:
    cover = {k: v for k, v in _context().cover.items() if k != "validity_date"}
    with pytest.raises(UndefinedError):
        render_document(load_template(), [], RenderContext(cover=cover))
