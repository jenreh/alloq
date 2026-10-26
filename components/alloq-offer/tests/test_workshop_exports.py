# ruff: noqa: ARG001
"""Snapshots, export requests, the export job and the download."""

import io
from datetime import UTC, datetime
from unittest.mock import AsyncMock, patch

import pytest
from alloq_offer.workshop.content.schema import Paragraph, ProseContent, Run
from alloq_offer.workshop.entities import (
    BlobEntity,
    ChapterEntity,
    DocState,
    ExportEntity,
    ExportKind,
    ExportStatus,
    JobEntity,
    SnapshotEntity,
    WorkspaceEntity,
)
from alloq_offer.workshop.export.docx_renderer import DRAFT_MARKER
from alloq_offer.workshop.export.models import ChapterInput, RenderError
from alloq_offer.workshop.export.verify import verify_document
from alloq_offer.workshop.jobs import export_docx
from alloq_offer.workshop.jobs.export_docx import ExportDocx
from alloq_offer.workshop.jobs.registry import JobContext, PermanentJobError
from alloq_offer.workshop.services import commercials, exports
from alloq_offer.workshop.services.chapters import ChapterService
from alloq_offer.workshop.services.exports import (
    EXPORT_JOB,
    ExportBlockedError,
    request_export,
)
from alloq_offer.workshop.services.outline import OutlineService
from alloq_offer.workshop.services.workspace import open_workspace
from alloq_offer.workshop.storage import DbBlobStore
from docx import Document
from sqlalchemy import select

NOW = datetime(2026, 10, 28, 9, 0, tzinfo=UTC)
KEY = "management_summary"


async def _prepared(db) -> None:  # noqa: ANN001
    async with db.session() as session:
        await open_workspace(session, 1, "Anna")
    service = ChapterService(db.session, 1, "Anna")
    await service.open()
    content = ProseContent(blocks=[Paragraph(runs=[Run(text="Fertiger Text")])])
    await service.save(KEY, 1, content)


async def _request(db, kind=ExportKind.DRAFT, **kwargs):  # noqa: ANN001, ANN003, ANN202
    options = {"confirmed": False, "unsaved": False} | kwargs
    return await request_export(
        db.session, 1, kind, user="Anna", now=NOW, max_attempts=3, **options
    )


async def _run_job(db, export_id: int) -> dict:  # noqa: ANN001
    store = DbBlobStore(10_000_000, db.session)
    ctx = JobContext(
        job_id=1, offer_id=1, kind=EXPORT_JOB, payload={"export_id": export_id}
    )
    return await ExportDocx(store, db.session).run(ctx)


async def test_draft_snapshot_is_fixed_and_job_queued(db) -> None:  # noqa: ANN001
    await _prepared(db)
    export = await _request(db)
    assert export.status == ExportStatus.WORKING
    assert export.filename.endswith("-Entwurf-1.docx")
    async with db.session() as session:
        snapshot = await session.scalar(select(SnapshotEntity))
        job = await session.scalar(select(JobEntity))
    assert snapshot is not None
    assert snapshot.number == 1
    chapter = next(c for c in snapshot.content["chapters"] if c["key"] == KEY)
    assert chapter["revision"] == 2
    assert job is not None
    assert job.kind == EXPORT_JOB
    assert job.payload == {"export_id": export.id}
    # Later edits do not change the snapshot.
    await ChapterService(db.session, 1, "Anna").save(
        KEY, 2, ProseContent(blocks=[Paragraph(runs=[Run(text="Neu")])])
    )
    async with db.session() as session:
        again = await session.scalar(select(SnapshotEntity))
    assert again is not None
    assert "Fertiger Text" in str(again.content)


async def test_export_follows_the_offer_outline(db) -> None:  # noqa: ANN001
    await _prepared(db)
    outline = OutlineService(db.session, 1, "Anna")
    await outline.move_chapter("verguetung", 0)
    await outline.remove_chapter("annahmen_risiken")
    added = await outline.add_chapter(
        "verguetung", chapter_type="prose", title="Datenschutz"
    )
    await outline.set_level(added, 2)
    export = await _request(db)
    async with db.session() as session:
        snapshot = await session.scalar(select(SnapshotEntity))
    assert snapshot is not None
    rows = [(c["key"], c["number"], c["title"]) for c in snapshot.content["chapters"]]
    assert rows[:3] == [
        ("verguetung", "1", "Vergütung"),
        (added, "1.1", "Datenschutz"),
        (KEY, "2", "Management Summary"),
    ]
    assert "annahmen_risiken" not in {row[0] for row in rows}
    await _run_job(db, export.id)
    async with db.session() as session:
        row = await session.get(ExportEntity, export.id)
        assert row is not None
        blob = await session.get(BlobEntity, row.blob_id)
    assert blob is not None
    headings = [
        (p.style.name, p.text)
        for p in Document(io.BytesIO(blob.data)).paragraphs
        if p.style is not None and p.style.name in {"Heading 1", "Heading 2"}
    ]
    titles = [text for _, text in headings]
    assert titles.index("Vergütung") < titles.index("Datenschutz")
    assert ("Heading 2", "Datenschutz") in headings
    assert "Annahmen und Risiken" not in titles
    # The frame of the layout stays after the chapters.
    assert titles.index("Rechtliche Bestimmungen") > titles.index("Laufzeit")


async def test_unsaved_changes_block_the_export(db) -> None:  # noqa: ANN001
    await _prepared(db)
    with pytest.raises(ExportBlockedError, match="ungespeicherte"):
        await _request(db, unsaved=True)


async def test_export_without_workspace_is_refused(db) -> None:  # noqa: ANN001
    with pytest.raises(ExportBlockedError):
        await _request(db)


async def test_final_is_blocked_with_all_reasons(db) -> None:  # noqa: ANN001
    await _prepared(db)
    with pytest.raises(ExportBlockedError) as caught:
        await _request(db, ExportKind.FINAL)
    reasons = caught.value.reasons
    assert any(r.startswith(exports.BLOCKER_PREFIX) for r in reasons)
    assert any("Kalkulation" in r for r in reasons)
    assert exports.NOT_CONFIRMED in reasons
    async with db.session() as session:
        assert await session.scalar(select(ExportEntity)) is None


async def test_final_passes_when_everything_is_decided(db) -> None:  # noqa: ANN001
    await _prepared(db)
    await commercials.confirm(db.session, 1, "Anna", NOW)
    with patch.object(exports.findings, "open_blockers", AsyncMock(return_value=[])):
        export = await _request(db, ExportKind.FINAL, confirmed=True)
    assert export.kind == ExportKind.FINAL


async def test_job_renders_stores_and_download_is_exact(db) -> None:  # noqa: ANN001
    await _prepared(db)
    export = await _request(db)
    result = await _run_job(db, export.id)
    assert result["status"] == "ok"
    async with db.session() as session:
        row = await session.get(ExportEntity, export.id)
    assert row is not None
    assert row.status == ExportStatus.READY
    assert row.manifest is not None
    assert row.manifest["chapters"][KEY] == 2
    assert row.manifest["sha256"] == row.sha256
    store = DbBlobStore(10_000_000, db.session)
    filename, data = await exports.load_download(db.session, store, 1, export.id)
    assert filename == row.filename
    assert len(data) == row.size
    document = Document(io.BytesIO(data))
    assert any(p.text == DRAFT_MARKER for p in document.paragraphs)
    assert any(p.text == "Fertiger Text" for p in document.paragraphs)


async def test_download_is_limited_to_the_offer_and_ready_files(db) -> None:  # noqa: ANN001
    await _prepared(db)
    export = await _request(db)
    store = DbBlobStore(10_000_000, db.session)
    with pytest.raises(ExportBlockedError):
        await exports.load_download(db.session, store, 1, export.id)
    await _run_job(db, export.id)
    with pytest.raises(ExportBlockedError):
        await exports.load_download(db.session, store, 2, export.id)
    with pytest.raises(ExportBlockedError):
        await exports.load_download(db.session, store, 1, 999)


async def test_download_refuses_a_changed_file(db) -> None:  # noqa: ANN001
    await _prepared(db)
    export = await _request(db)
    await _run_job(db, export.id)
    async with db.session() as session:
        row = await session.get(ExportEntity, export.id)
        assert row is not None
        row.sha256 = "0" * 64
    store = DbBlobStore(10_000_000, db.session)
    with pytest.raises(ExportBlockedError, match="Prüfsumme"):
        await exports.load_download(db.session, store, 1, export.id)


async def test_final_job_marks_the_document_final(db) -> None:  # noqa: ANN001
    await _prepared(db)
    await commercials.confirm(db.session, 1, "Anna", NOW)
    with patch.object(exports.findings, "open_blockers", AsyncMock(return_value=[])):
        export = await _request(db, ExportKind.FINAL, confirmed=True)
    await _run_job(db, export.id)
    store = DbBlobStore(10_000_000, db.session)
    _, data = await exports.load_download(db.session, store, 1, export.id)
    assert DRAFT_MARKER not in [p.text for p in Document(io.BytesIO(data)).paragraphs]
    async with db.session() as session:
        workspace = await session.scalar(select(WorkspaceEntity))
    assert workspace is not None
    assert workspace.doc_state == DocState.FINAL


async def test_draft_lists_open_points(db) -> None:  # noqa: ANN001
    await _prepared(db)
    export = await _request(db)
    await _run_job(db, export.id)
    store = DbBlobStore(10_000_000, db.session)
    _, data = await exports.load_download(db.session, store, 1, export.id)
    texts = [p.text for p in Document(io.BytesIO(data)).paragraphs]
    assert "Offene Punkte" in texts


async def test_a_failed_render_is_not_shown_as_success(db) -> None:  # noqa: ANN001
    await _prepared(db)
    export = await _request(db)
    with (
        patch.object(export_docx, "render_document", side_effect=RenderError("x")),
        pytest.raises(PermanentJobError),
    ):
        await _run_job(db, export.id)
    async with db.session() as session:
        row = await session.get(ExportEntity, export.id)
    assert row is not None
    assert row.status == ExportStatus.FAILED
    assert row.error == export_docx.FAILED
    assert row.blob_id is None
    # A new request is possible and starts clean.
    again = await _request(db)
    assert again.id != export.id


async def test_exhausted_job_marks_the_export_failed(db) -> None:  # noqa: ANN001
    await _prepared(db)
    export = await _request(db)
    ctx = JobContext(1, 1, EXPORT_JOB, payload={"export_id": export.id})
    with patch.object(export_docx, "get_asyncdb_session", db.session):
        await export_docx._on_exhausted(ctx)  # noqa: SLF001
    rows = await exports.list_exports(db.session, 1)
    assert rows[0]["failed"]


async def test_missing_export_is_a_permanent_error(db) -> None:  # noqa: ANN001
    with pytest.raises(PermanentJobError):
        await _run_job(db, 99)


async def test_template_version_binding(db) -> None:  # noqa: ANN001
    await _prepared(db)
    async with db.session() as session:
        workspace = await session.scalar(select(WorkspaceEntity))
        assert workspace is not None
        workspace.template_version = "0.0-old"
    with pytest.raises(ExportBlockedError, match="Vorlage"):
        await _request(db)


async def test_list_shows_newest_first(db) -> None:  # noqa: ANN001
    await _prepared(db)
    first = await _request(db)
    second = await _request(db)
    rows = await exports.list_exports(db.session, 1)
    assert [r["id"] for r in rows] == [second.id, first.id]


def test_verify_rejects_missing_chapter_and_wrong_marker() -> None:
    buffer = io.BytesIO()
    document = Document()
    document.add_heading("Eins", level=1)
    document.save(buffer)
    content = ProseContent()
    chapters = [ChapterInput("a", "Eins", 1, content)]
    verify_document(buffer.getvalue(), chapters, draft=False)
    with pytest.raises(RenderError):
        verify_document(buffer.getvalue(), chapters, draft=True)
    with pytest.raises(RenderError, match="Zwei"):
        verify_document(
            buffer.getvalue(), [ChapterInput("b", "Zwei", 1, content)], draft=False
        )
    with pytest.raises(RenderError):
        verify_document(b"not a docx", chapters, draft=False)


async def test_final_needs_required_chapters_to_be_reviewed(db) -> None:  # noqa: ANN001
    await _prepared(db)
    await commercials.confirm(db.session, 1, "Anna", NOW)
    async with db.session() as session:
        chapter = await session.scalar(
            select(ChapterEntity).where(ChapterEntity.chapter_key == KEY)
        )
        assert chapter is not None
        chapter.needs_review = True
    with (
        patch.object(exports.findings, "open_blockers", AsyncMock(return_value=[])),
        pytest.raises(ExportBlockedError) as caught,
    ):
        await _request(db, ExportKind.FINAL, confirmed=True)
    assert caught.value.reasons == [f"{exports.UNREVIEWED_PREFIX}Management Summary"]


def test_stories_stored_with_the_old_scope_still_load() -> None:
    stored = {
        "key": "US-01",
        "title": "T",
        "sentence": "Als A möchte ich b, um c.",
        "scope": "angeboten",
        "group": "G",
    }
    story = export_docx._story(stored)  # noqa: SLF001
    assert (story.key, story.group) == ("US-01", "G")
