"""Tests of the source parsers."""

import io
from email.message import EmailMessage

import httpx
import pytest
from alloq_offer.workshop.configuration import AzureOpenAiConfig
from alloq_offer.workshop.extraction import (
    ExtractionError,
    parse_docx,
    parse_eml,
    parse_html,
    parse_pdf,
    parse_pptx,
    parse_text,
    to_result,
)
from alloq_offer.workshop.extraction.audio import (
    AUDIO_TOO_LARGE,
    NO_TRANSCRIBER,
    AzureTranscriber,
    TranscriptionUnavailableError,
    TranscriptSegment,
)
from alloq_offer.workshop.extraction.models import KIND_QUOTE, KIND_SIGNATURE
from alloq_offer.workshop.extraction.pdf import NO_TEXT_LAYER
from alloq_offer.workshop.extraction.slides import IMAGE_NOT_READ
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)

PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d49444154789c6360f8cfc000000301010018dd8db00000000049454e44ae426082"
)


def _text_pdf(*texts: str) -> bytes:
    writer = PdfWriter()
    font = writer._add_object(
        DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
    )
    for text in texts:
        page = writer.add_blank_page(200, 200)
        if text:
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode())
            page[NameObject("/Contents")] = writer._add_object(stream)
            page[NameObject("/Resources")] = DictionaryObject(
                {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
            )
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_text_splits_paragraphs_and_numbers_them() -> None:
    result = parse_text(b"Erster Absatz.\n\n\nZweiter Absatz.\r\nNoch eine Zeile.")
    assert [block.text for block in result.blocks] == [
        "Erster Absatz.",
        "Zweiter Absatz.\nNoch eine Zeile.",
    ]
    assert result.blocks[1].anchor == {"paragraph": 2}


def test_text_reads_legacy_encoding() -> None:
    assert parse_text("Größe".encode("cp1252")).blocks[0].text == "Größe"


def test_text_without_content_has_no_blocks() -> None:
    assert parse_text(b" \n\n ").blocks == []


def _eml(body: str, *, attachment: bytes | None = None) -> bytes:
    message = EmailMessage()
    message["Subject"] = "Anfrageprozess"
    message["From"] = "anna@example.com"
    message["Date"] = "Mon, 05 Oct 2026 09:00:00 +0000"
    message["Message-ID"] = "<abc@example.com>"
    message.set_content(body)
    if attachment is not None:
        message.add_attachment(
            attachment, maintype="application", subtype="pdf", filename="Konzept.pdf"
        )
    return message.as_bytes()


def test_eml_separates_body_quote_and_signature() -> None:
    result = parse_eml(
        _eml("Hallo Team,\n\nwir brauchen eine Lösung.\n\n> alte Frage\n\n-- \nAnna")
    )
    kinds = {block.kind for block in result.blocks}
    assert {KIND_QUOTE, KIND_SIGNATURE} <= kinds
    assert result.title == "Anfrageprozess"
    header = result.blocks[0]
    assert "Betreff: Anfrageprozess" in header.text
    assert header.anchor["message_id"] == "<abc@example.com>"
    assert any("Lösung" in block.text for block in result.blocks)


def test_eml_attachments_become_attachments() -> None:
    result = parse_eml(_eml("Siehe Anhang.", attachment=b"%PDF-1.4"))
    assert [(a.filename, a.mime, a.data) for a in result.attachments] == [
        ("Konzept.pdf", "application/pdf", b"%PDF-1.4")
    ]


def test_eml_html_only_body_is_stripped_of_markup() -> None:
    message = EmailMessage()
    message["Subject"] = "Html"
    message.set_content("<p>Hallo <b>Welt</b></p>", subtype="html")
    result = parse_eml(message.as_bytes())
    assert any(
        "Hallo" in block.text and "<" not in block.text for block in result.blocks
    )


def test_docx_keeps_document_order_of_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("Einleitung")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Rolle"
    table.cell(0, 1).text = "Aufgabe"
    table.cell(1, 0).text = "Leitung"
    table.cell(1, 1).text = "Freigabe"
    document.add_paragraph("")
    document.add_paragraph("Schluss")
    buffer = io.BytesIO()
    document.save(buffer)

    result = parse_docx(buffer.getvalue())

    assert [(b.kind, b.anchor) for b in result.blocks] == [
        ("paragraph", {"paragraph": 1}),
        ("table", {"table": 1}),
        ("paragraph", {"paragraph": 3}),
    ]
    assert result.blocks[1].text == "Rolle | Aufgabe\nLeitung | Freigabe"


def test_docx_rejects_garbage() -> None:
    with pytest.raises(ExtractionError):
        parse_docx(b"not a zip")


def test_pdf_reads_text_pages_and_flags_pages_without_text_layer() -> None:
    result = parse_pdf(_text_pdf("Hallo Welt", ""))
    assert [b.text for b in result.text_blocks] == ["Hallo Welt"]
    assert result.gaps[0].text == NO_TEXT_LAYER
    assert result.gaps[0].anchor == {"page": 2}


def test_pdf_rejects_garbage_and_encrypted_files() -> None:
    with pytest.raises(ExtractionError):
        parse_pdf(b"%PDF-broken")
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    writer.encrypt("geheim")
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(ExtractionError, match="passwortgeschützt"):
        parse_pdf(buffer.getvalue())


def test_pptx_reads_text_tables_notes_and_flags_pictures() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Zielbild"
    slide.shapes.add_picture(io.BytesIO(PNG), Inches(1), Inches(2))
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1))
    table.table.cell(0, 0).text = "A"
    table.table.cell(0, 1).text = "B"
    slide.notes_slide.notes_text_frame.text = "Sprechernotiz"
    buffer = io.BytesIO()
    presentation.save(buffer)

    result = parse_pptx(buffer.getvalue())

    assert [b.text for b in result.text_blocks] == [
        "Zielbild",
        "A | B",
        "Sprechernotiz",
    ]
    assert result.gaps[0].text == IMAGE_NOT_READ
    assert result.gaps[0].anchor["slide"] == 1


def test_pptx_rejects_garbage() -> None:
    with pytest.raises(ExtractionError):
        parse_pptx(b"nope")


def test_transcript_uses_technical_speaker_labels() -> None:
    result = to_result(
        [
            TranscriptSegment(0, 1500, "Guten Tag", "A"),
            TranscriptSegment(1500, 3000, "Hallo", "B"),
            TranscriptSegment(3000, 4000, "  ", "A"),
            TranscriptSegment(4000, 5000, "Weiter", "A"),
        ]
    )
    assert [(b.speaker, b.text) for b in result.blocks] == [
        ("Sprecher 1", "Guten Tag"),
        ("Sprecher 2", "Hallo"),
        ("Sprecher 1", "Weiter"),
    ]
    assert result.blocks[0].anchor == {"start_ms": 0, "end_ms": 1500}


async def test_azure_transcriber_without_deployment_is_unavailable() -> None:
    transcriber = AzureTranscriber(AzureOpenAiConfig())
    with pytest.raises(TranscriptionUnavailableError, match=NO_TRANSCRIBER):
        await transcriber.transcribe(b"x", "a.wav", "audio/wav")


async def test_azure_transcriber_rejects_files_above_the_limit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr("alloq_offer.workshop.extraction.audio.MAX_AUDIO_BYTES", 3)
    config = AzureOpenAiConfig(endpoint="https://x", transcription_deployment="t")
    with pytest.raises(TranscriptionUnavailableError, match=AUDIO_TOO_LARGE):
        await AzureTranscriber(config).transcribe(b"abcd", "a.wav", "audio/wav")


async def test_azure_transcriber_posts_audio_and_parses_segments(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    seen: dict = {}

    def handler(request: httpx.Request) -> httpx.Response:
        seen["url"] = str(request.url)
        seen["key"] = request.headers["api-key"]
        seen["body"] = request.content
        return httpx.Response(
            200,
            json={
                "segments": [
                    {"start": 0.5, "end": 2.0, "text": "Hallo", "speaker": "A"}
                ]
            },
        )

    real_client = httpx.AsyncClient
    monkeypatch.setattr(
        "alloq_offer.workshop.extraction.audio.httpx.AsyncClient",
        lambda **kw: real_client(transport=httpx.MockTransport(handler), **kw),
    )
    config = AzureOpenAiConfig(
        endpoint="https://x.openai.azure.com/",
        api_key="secret",  # ty: ignore[invalid-argument-type]
        transcription_deployment="gpt-4o-transcribe-diarize",
    )

    segments = await AzureTranscriber(config).transcribe(b"audio", "a.wav", "audio/wav")

    assert segments == [TranscriptSegment(500, 2000, "Hallo", "A")]
    assert "/deployments/gpt-4o-transcribe-diarize/audio/transcriptions" in seen["url"]
    assert seen["key"] == "secret"
    assert b"diarized_json" in seen["body"]


def test_html_keeps_text_and_drops_markup_scripts_and_styles() -> None:
    html = (
        "<html><head><title>Titel</title><style>p{}</style></head><body>"
        "<h1>Anfrage</h1><p>Wir brauchen <b>Hilfe</b>.</p><script>x()</script>"
        "<p>Zweiter &amp; letzter Absatz</p></body></html>"
    )
    result = parse_html(html.encode())
    assert [b.text for b in result.blocks] == [
        "Anfrage",
        "Wir brauchen Hilfe.",
        "Zweiter & letzter Absatz",
    ]
