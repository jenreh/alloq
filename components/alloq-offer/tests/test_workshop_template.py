"""Template packages: published versions from the database and the Word layout."""

import io
import zipfile

import pytest
from alloq_offer.workshop.entities import BlobEntity
from alloq_offer.workshop.export.template_package import (
    CHAPTER_ANCHOR,
    NO_ANCHOR,
    ChapterDef,
    TemplateError,
    TemplatePackage,
    check_layout,
    clean_layout,
    frame_headings,
)
from alloq_offer.workshop.services import templates
from alloq_offer.workshop.template_entities import OfferTemplateVersionEntity
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pydantic import ValidationError
from sqlalchemy import select

ATTACHED_TEMPLATE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    "attachedTemplate"
)


def _docx(*lines: str) -> bytes:
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_seeded_package_has_every_chapter_type(quo_v1: TemplatePackage) -> None:
    keys = [c.key for c in quo_v1.chapters]
    assert keys[0] == "management_summary"
    assert {c.chapter_type for c in quo_v1.chapters} == {
        "prose",
        "table",
        "stories",
        "calc",
    }
    assert quo_v1.chapters[0].settings().required_info == [
        "Ziel des Kunden",
        "Angebotene Lösung",
        "Umfang",
    ]


def test_chapter_start_content_matches_its_type(quo_v1: TemplatePackage) -> None:
    for chapter in quo_v1.chapters:
        assert chapter.content().chapter_type == chapter.chapter_type
    with pytest.raises(ValidationError):
        ChapterDef(
            key="x",
            title="X",
            chapter_type="prose",
            default_content={"chapter_type": "table", "columns": ["A"]},
        )


def test_layout_has_single_anchor_and_no_customer_data(
    quo_v1: TemplatePackage,
) -> None:
    quo_v1.check_layout()
    text = "\n".join(p.text for p in check_layout(quo_v1.layout_bytes).paragraphs)
    assert text.count(CHAPTER_ANCHOR) == 1
    assert "Treysta" not in text.title()


def test_seeded_frame_matches_the_layout(quo_v1: TemplatePackage) -> None:
    assert frame_headings(quo_v1.layout_bytes) == (
        quo_v1.frame_before,
        quo_v1.frame_after,
    )


def test_duplicate_chapter_keys_rejected(quo_v1: TemplatePackage) -> None:
    chapter = quo_v1.chapters[0].model_dump()
    with pytest.raises(ValidationError, match="eindeutig"):
        TemplatePackage(
            key="x", version="1", chapters=[chapter, chapter], layout_bytes=b""
        )


def test_layout_without_anchor_or_not_word_is_rejected() -> None:
    with pytest.raises(TemplateError, match="Anker"):
        TemplatePackage(
            key="x", version="1", chapters=[], layout_bytes=_docx("Text")
        ).check_layout()
    with pytest.raises(TemplateError, match="Word"):
        check_layout(b"kein Word")
    assert NO_ANCHOR.endswith(f"{CHAPTER_ANCHOR}.")


def test_frame_headings_split_at_the_anchor() -> None:
    document = Document()
    document.add_heading("Vorne", 1)
    document.add_paragraph("Fließtext")
    document.add_paragraph(CHAPTER_ANCHOR)
    document.add_heading("Hinten", 1)
    out = io.BytesIO()
    document.save(out)
    assert frame_headings(out.getvalue()) == (["Vorne"], ["Hinten"])


def test_clean_layout_drops_author_and_keeps_the_anchor(
    quo_v1: TemplatePackage,
) -> None:
    document = Document(io.BytesIO(_docx(CHAPTER_ANCHOR)))
    document.core_properties.author = "Erika Muster"
    out = io.BytesIO()
    document.save(out)
    cleaned = clean_layout(out.getvalue())
    assert Document(io.BytesIO(cleaned)).core_properties.author == ""
    check_layout(cleaned)
    names = zipfile.ZipFile(io.BytesIO(clean_layout(quo_v1.layout_bytes))).namelist()
    assert not [n for n in names if "comments" in n or "people" in n]


def test_clean_layout_drops_tracked_changes_and_linked_files() -> None:
    document = Document(io.BytesIO(_docx(CHAPTER_ANCHOR, "Alt")))
    paragraph = document.paragraphs[1]._p  # noqa: SLF001
    run = paragraph.r_lst[0]
    deleted = OxmlElement("w:del")
    inserted = OxmlElement("w:ins")
    paragraph.append(deleted)
    paragraph.append(inserted)
    inserted.append(run)
    document.part.relate_to(
        "file://server/vorlage.dotm", ATTACHED_TEMPLATE, is_external=True
    )
    document.part.relate_to("https://example.org", RT.HYPERLINK, is_external=True)
    out = io.BytesIO()
    document.save(out)
    cleaned = Document(io.BytesIO(clean_layout(out.getvalue())))
    body = cleaned.element.body
    assert not list(body.iter(qn("w:del")))
    assert not list(body.iter(qn("w:ins")))
    assert cleaned.paragraphs[1].text == "Alt"
    targets = {rel.target_ref for rel in cleaned.part.rels.values() if rel.is_external}
    assert targets == {"https://example.org"}


def test_version_number_maps_legacy_texts() -> None:
    assert templates.version_number("quo_v1.0") == 1
    assert templates.version_number("3") == 3
    with pytest.raises(TemplateError, match=templates.UNKNOWN_VERSION):
        templates.version_number("quo_v9")


async def test_loaded_layout_is_cleaned(db) -> None:
    document = Document(io.BytesIO(_docx(CHAPTER_ANCHOR)))
    document.core_properties.author = "Erika Muster"
    out = io.BytesIO()
    document.save(out)
    async with db.session() as session:
        blob = await session.get(BlobEntity, 1)
        assert blob is not None
        blob.data = out.getvalue()
    async with db.session() as session:
        package = await templates.load_template(session, "quo_v1", 1)
    assert Document(io.BytesIO(package.layout_bytes)).core_properties.author == ""


async def test_export_check_needs_a_published_version(db) -> None:
    async with db.session() as session:
        assert await templates.template_exists(session, "quo_v1", "quo_v1.0")
        assert not await templates.template_exists(session, "quo_v1", "2")
        assert not await templates.template_exists(session, "quo_v1", "alt")


async def test_load_by_version_caches_and_rejects_unknown(db) -> None:
    async with db.session() as session:
        legacy = await templates.load_template(session, "quo_v1", "quo_v1.0")
        latest = await templates.latest_template(session, "quo_v1")
        assert legacy is latest
        assert legacy.version == "1"
        assert legacy.frame_after[-1] == "Unterschriften"
        with pytest.raises(TemplateError, match=templates.UNKNOWN_VERSION):
            await templates.load_template(session, "quo_v1", 2)
        with pytest.raises(TemplateError, match=templates.NO_VERSION):
            await templates.latest_template(session, "gibt_es_nicht")
        # Versions never change, so a later read is served from the cache.
        row = await session.scalar(select(OfferTemplateVersionEntity))
        assert row is not None
        row.payload = {"chapters": [{"key": "ohne_typ"}]}
    async with db.session() as session:
        cached = await templates.load_template(session, "quo_v1", 1)
    assert cached is legacy
    templates.clear_cache()
    with pytest.raises(TemplateError, match=templates.BROKEN_VERSION):
        async with db.session() as session:
            await templates.load_template(session, "quo_v1", 1)
