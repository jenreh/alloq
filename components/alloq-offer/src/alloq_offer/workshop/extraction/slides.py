"""PowerPoint (PPTX): text frames, tables and notes per slide."""

import io

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape

from alloq_offer.workshop.extraction.models import (
    KIND_NOTE,
    KIND_PARAGRAPH,
    KIND_TABLE,
    ExtractedBlock,
    ExtractionError,
    ExtractionResult,
    gap,
)

IMAGE_NOT_READ = "Bildinhalt nicht ausgewertet"


def _shape_blocks(
    shape: BaseShape, slide_no: int, shape_no: int
) -> list[ExtractedBlock]:
    anchor = {"slide": slide_no, "shape": shape_no}
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return [gap(IMAGE_NOT_READ, anchor)]
    if isinstance(shape, GroupShape):
        return [
            block
            for index, child in enumerate(shape.shapes, 1)
            for block in _shape_blocks(child, slide_no, shape_no * 100 + index)
        ]
    if isinstance(shape, GraphicFrame) and shape.has_table:
        rows = (
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in shape.table.rows
        )
        return [ExtractedBlock(KIND_TABLE, "\n".join(rows), anchor)]
    if isinstance(shape, Shape) and shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text:
            return [ExtractedBlock(KIND_PARAGRAPH, text, anchor)]
    return []


def parse_pptx(data: bytes) -> ExtractionResult:
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - python-pptx raises many types
        raise ExtractionError("Die Präsentation konnte nicht gelesen werden.") from exc
    blocks: list[ExtractedBlock] = []
    for slide_no, slide in enumerate(presentation.slides, 1):
        for shape_no, shape in enumerate(slide.shapes, 1):
            blocks.extend(_shape_blocks(shape, slide_no, shape_no))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                blocks.append(
                    ExtractedBlock(KIND_NOTE, notes, {"slide": slide_no, "note": 1})
                )
    return ExtractionResult(blocks=blocks)
