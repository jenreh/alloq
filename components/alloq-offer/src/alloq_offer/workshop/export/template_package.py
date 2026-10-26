"""Template package: chapter outline, chapter settings and Word layout of a version."""

from __future__ import annotations

import io
from typing import Any, Literal

from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.package import OpcPackage
from docx.opc.part import Part, XmlPart
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.text.paragraph import Paragraph as DocxParagraph
from pydantic import BaseModel, Field, model_validator

from alloq_offer.workshop.content.schema import parse_content
from alloq_offer.workshop.content.text import Content

DEFAULT_TEMPLATE_KEY = "quo_v1"
CHAPTER_ANCHOR = "[[KAPITEL]]"
NO_ANCHOR = f"Die Vorlage braucht genau einen Anker {CHAPTER_ANCHOR}."
BROKEN_LAYOUT = "Die Datei ist kein lesbares Word-Dokument."
ChapterType = Literal["prose", "table", "stories", "calc"]
# Parts that carry reviewer comments, people, labels, properties or statistics.
PRIVATE_PARTS = (
    "/comments",
    "/commentsIds",
    "/commentsExtended",
    "/commentsExtensible",
    "/people",
    "/customXml",
    "/custom-properties",
    "/extended-properties",
    "/thumbnail",
)
HYPERLINK = "/hyperlink"
PRIVATE_PROPERTIES = (
    "author",
    "last_modified_by",
    "comments",
    "keywords",
    "category",
    "subject",
)
# Comment anchors and tracked changes with their authors: deleted and moved-away
# text and the old formatting go; inserted and moved-in text stays (KEPT_MARKUP).
DROPPED_MARKUP = (
    "w:commentRangeStart",
    "w:commentRangeEnd",
    "w:commentReference",
    "w:del",
    "w:moveFrom",
    "w:moveFromRangeStart",
    "w:moveFromRangeEnd",
    "w:moveToRangeStart",
    "w:moveToRangeEnd",
    "w:rPrChange",
    "w:pPrChange",
    "w:sectPrChange",
    "w:tblPrChange",
    "w:tblGridChange",
    "w:trPrChange",
    "w:tcPrChange",
    "w:numberingChange",
)
KEPT_MARKUP = ("w:ins", "w:moveTo")


class TemplateError(ValueError):
    """The template package is inconsistent."""


class ChapterSettings(BaseModel):
    """What the AI and the checks know about one chapter."""

    guiding_question: str = ""
    content_rules: list[str] = Field(default_factory=list)
    prompt_text: str = ""
    required_info: list[str] = Field(default_factory=list)
    checks: list[str] = Field(default_factory=list)


SETTINGS_FIELDS = frozenset(ChapterSettings.model_fields)


class ChapterDef(ChapterSettings):
    """One chapter of a published template version, settings resolved."""

    key: str
    title: str
    level: Literal[1, 2] = 1
    required: bool = False
    chapter_type: ChapterType
    default_content: dict[str, Any] = Field(default_factory=dict)

    @model_validator(mode="after")
    def _content_matches_type(self) -> ChapterDef:
        if self.content().chapter_type != self.chapter_type:
            raise TemplateError("Der Inhalt passt nicht zum Kapiteltyp.")
        return self

    def content(self) -> Content:
        """The content a new chapter starts with."""
        return parse_content(
            {"chapter_type": self.chapter_type, **self.default_content}
        )

    def settings(self) -> ChapterSettings:
        return ChapterSettings.model_validate(self.model_dump(include=SETTINGS_FIELDS))


class TemplatePackage(BaseModel):
    key: str
    version: str
    chapters: list[ChapterDef]
    frame_before: list[str] = Field(default_factory=list)
    frame_after: list[str] = Field(default_factory=list)
    layout_bytes: bytes

    @model_validator(mode="after")
    def _unique_keys(self) -> TemplatePackage:
        keys = [chapter.key for chapter in self.chapters]
        if len(keys) != len(set(keys)):
            raise TemplateError("Kapitelschlüssel müssen eindeutig sein.")
        return self

    def check_layout(self) -> None:
        """Fail unless the layout holds the chapter anchor exactly once."""
        check_layout(self.layout_bytes)


def _document(layout: bytes) -> DocxDocument:
    try:
        return Document(io.BytesIO(layout))
    except Exception as exc:  # python-docx raises several unrelated types
        raise TemplateError(BROKEN_LAYOUT) from exc


def check_layout(layout: bytes) -> DocxDocument:
    """The parsed layout; it must hold the chapter anchor exactly once."""
    document = _document(layout)
    anchors = [p for p in document.paragraphs if p.text.strip() == CHAPTER_ANCHOR]
    if len(anchors) != 1:
        raise TemplateError(NO_ANCHOR)
    return document


def _drop_private_rels(owner: Part | OpcPackage) -> None:
    for r_id, rel in list(owner.rels.items()):
        linked = rel.is_external and not rel.reltype.endswith(HYPERLINK)
        if linked or rel.reltype.endswith(PRIVATE_PARTS):
            del owner.rels[r_id]


def _strip_markup(element: BaseOxmlElement) -> None:
    for name in DROPPED_MARKUP:
        for node in list(element.iter(qn(name))):
            node.getparent().remove(node)
    for name in KEPT_MARKUP:
        for wrapper in list(element.iter(qn(name))):
            for child in list(wrapper):
                wrapper.addprevious(child)
            wrapper.getparent().remove(wrapper)


def clean_layout(layout: bytes) -> bytes:
    """The layout without comments, tracked changes, labels, links and authors.

    Every export is rendered from the layout, so nothing internal may stay in it.
    """
    document = check_layout(layout)
    package = document.part.package
    parts = list(package.iter_parts())
    _drop_private_rels(package)
    for part in parts:
        _drop_private_rels(part)
        if isinstance(part, XmlPart):
            _strip_markup(part.element)
    properties = document.core_properties
    for name in PRIVATE_PROPERTIES:
        setattr(properties, name, "")
    properties.revision = 1
    printed = properties._element.find(qn("cp:lastPrinted"))  # noqa: SLF001
    if printed is not None:
        properties._element.remove(printed)  # noqa: SLF001
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _is_heading(paragraph: DocxParagraph) -> bool:
    style = (paragraph.style.name if paragraph.style is not None else "").lower()
    return bool(paragraph.text.strip()) and (
        style.startswith("heading") or "überschrift" in style
    )


def frame_headings(layout: bytes) -> tuple[list[str], list[str]]:
    """Headings of the fixed frame before and after the chapter anchor."""
    before: list[str] = []
    after: list[str] = []
    target = before
    for paragraph in check_layout(layout).paragraphs:
        if paragraph.text.strip() == CHAPTER_ANCHOR:
            target = after
        elif _is_heading(paragraph):
            target.append(paragraph.text.strip())
    return before, after


def find_anchor(document: DocxDocument) -> DocxParagraph:
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == CHAPTER_ANCHOR:
            return paragraph
    raise TemplateError(NO_ANCHOR)
