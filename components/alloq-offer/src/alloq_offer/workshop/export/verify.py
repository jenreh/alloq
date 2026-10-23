"""Structural check of a rendered offer document (spec 13.1 step 5)."""

import io
from collections.abc import Sequence

from docx import Document

from alloq_offer.workshop.export.docx_renderer import DRAFT_MARKER
from alloq_offer.workshop.export.models import ChapterInput, RenderError

BROKEN = "Die erzeugte Datei ist nicht vollständig."


def verify_document(
    data: bytes, chapters: Sequence[ChapterInput], *, draft: bool
) -> None:
    """Fail unless the chapters appear in order and the draft marker fits the kind."""
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise RenderError(BROKEN) from exc
    headings = [
        p.text
        for p in document.paragraphs
        if p.style is not None and p.style.name.startswith("Heading")
    ]
    position = 0
    for chapter in chapters:
        try:
            position = headings.index(chapter.title, position) + 1
        except ValueError as exc:
            raise RenderError(f"{BROKEN} Kapitel fehlt: {chapter.title}") from exc
    marked = any(p.text == DRAFT_MARKER for p in document.paragraphs)
    if marked != draft:
        raise RenderError(BROKEN)
