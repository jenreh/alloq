"""One renderer per chapter type, selected through a registry."""

from __future__ import annotations

import re
from collections.abc import Callable

from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph as DocxParagraph

from alloq_offer.workshop.content.schema import (
    BOUND_FIELDS,
    SCOPE_LABELS,
    BoundValue,
    BulletList,
    CalcContent,
    Heading,
    Inline,
    OrderedList,
    Paragraph,
    ProseContent,
    Run,
    StoriesContent,
    TableBlock,
    TableContent,
    plain_text,
)
from alloq_offer.workshop.export.models import (
    NO_GROUP_LABEL,
    ChapterInput,
    RenderContext,
    RenderError,
    ResolvedStory,
    ResolvedValue,
    UnsupportedContentError,
)

TABLE_STYLE = "arvato-Tabelle blau"
BULLET_STYLE = "Aufzählung 1"
ORDERED_STYLE = "List Number"
NO_STORIES_TEXT = "Keine Stories im gewählten Umfang."
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class Writer:
    """Appends content in front of the chapter anchor of a python-docx document."""

    def __init__(self, document: DocxDocument, anchor: DocxParagraph) -> None:
        self._doc = document
        self._anchor = anchor

    def _style(self, name: str) -> str:
        try:
            return self._doc.styles[name].name
        except KeyError as exc:
            raise RenderError(f"Formatvorlage fehlt in der Vorlage: {name}") from exc

    def _place(self, element: object) -> None:
        self._anchor._p.addprevious(element)  # noqa: SLF001

    def runs(self, paragraph: DocxParagraph, runs: Inline) -> None:
        for run in runs:
            target = paragraph.add_run(_ILLEGAL_XML.sub("", run.text))
            target.bold = run.bold or None
            target.italic = run.italic or None

    def paragraph(self, runs: Inline, style: str | None = None) -> DocxParagraph:
        paragraph = self._anchor.insert_paragraph_before(
            style=self._style(style) if style else None
        )
        self.runs(paragraph, runs)
        return paragraph

    def heading(self, text: str, level: int) -> None:
        self._anchor.insert_paragraph_before(
            _ILLEGAL_XML.sub("", text), style=self._style(f"Heading {level}")
        )

    def table(self, header: list[Inline], rows: list[list[Inline]]) -> None:
        table = self._doc.add_table(rows=1 + len(rows), cols=len(header))
        table.style = self._doc.styles[self._style(TABLE_STYLE)]
        for index, cell_runs in enumerate(header):
            self.runs(table.rows[0].cells[index].paragraphs[0], cell_runs)
        for row_index, row in enumerate(rows, start=1):
            for index, cell_runs in enumerate(row):
                self.runs(table.rows[row_index].cells[index].paragraphs[0], cell_runs)
        tr_pr = table.rows[0]._tr.get_or_add_trPr()  # noqa: SLF001
        tr_pr.append(tr_pr.makeelement(qn("w:tblHeader"), {}))
        self._place(table._tbl)  # noqa: SLF001


Renderer = Callable[[Writer, ChapterInput, RenderContext], None]


def _text(value: str) -> Inline:
    return [Run(text=value)]


def _bold(value: str) -> Inline:
    return [Run(text=value, bold=True)]


def _resolve(ctx: RenderContext, field: str) -> ResolvedValue:
    return ctx.values.get(field, ResolvedValue(text=None))


def _render_prose(writer: Writer, chapter: ChapterInput, ctx: RenderContext) -> None:
    content = chapter.content
    if not isinstance(content, ProseContent):
        raise UnsupportedContentError(chapter.key)
    inner_min = min(chapter.level + 1, 4)
    for block in content.blocks:
        match block:
            case Paragraph():
                writer.paragraph(block.runs)
            case Heading():
                writer.heading(plain_text(block.runs), max(block.level, inner_min))
            case BulletList():
                for item in block.items:
                    writer.paragraph(item, BULLET_STYLE)
            case OrderedList():
                for item in block.items:
                    writer.paragraph(item, ORDERED_STYLE)
            case TableBlock():
                writer.table(block.header, block.rows)
            case BoundValue():
                label = BOUND_FIELDS[block.field]
                writer.paragraph(
                    _text(f"{label}: {_resolve(ctx, block.field).display()}")
                )
            case _:
                raise UnsupportedContentError(type(block).__name__)


def _render_table(writer: Writer, chapter: ChapterInput, _ctx: RenderContext) -> None:
    content = chapter.content
    if not isinstance(content, TableContent):
        raise UnsupportedContentError(chapter.key)
    writer.table([_bold(c) for c in content.columns], content.rows)


def _render_stories(writer: Writer, chapter: ChapterInput, ctx: RenderContext) -> None:
    content = chapter.content
    if not isinstance(content, StoriesContent):
        raise UnsupportedContentError(chapter.key)
    stories = [s for s in ctx.stories if s.scope in content.scopes]
    if not stories:
        writer.paragraph(_text(NO_STORIES_TEXT))
        return
    heading_level = min(chapter.level + 1, 4)
    groups: dict[str, list[ResolvedStory]] = {}
    for story in stories:
        groups.setdefault(story.group or NO_GROUP_LABEL, []).append(story)
    for group, items in groups.items():
        writer.heading(group, heading_level)
        for story in items:
            label = SCOPE_LABELS.get(story.scope, story.scope)
            head = _bold(f"{story.key} – {story.title}")
            head.extend(_text(f" ({label})"))
            writer.paragraph(head)
            writer.paragraph(_text(story.sentence))


def _render_calc(writer: Writer, chapter: ChapterInput, ctx: RenderContext) -> None:
    content = chapter.content
    if not isinstance(content, CalcContent):
        raise UnsupportedContentError(chapter.key)
    if content.intro:
        writer.paragraph(content.intro)
    rows = [
        [_text(row.label), _text(_resolve(ctx, row.field).display())]
        for row in content.rows
    ]
    writer.table([_bold("Position"), _bold("Wert")], rows)


RENDERERS: dict[str, Renderer] = {
    "prose": _render_prose,
    "table": _render_table,
    "stories": _render_stories,
    "calc": _render_calc,
}
