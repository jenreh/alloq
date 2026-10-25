# ruff: noqa: E501
"""Create a complete demo offer in the workshop of the configured database.

Adds an offer with five customer sources (Word, PDF, PowerPoint, e-mail, pasted
text), runs extraction, the one-click draft (map, briefing, chapters, review,
Word draft). Only jobs of the new offer are run here.

Needs ``ai_provider: demo`` (local profile) or a configured Azure OpenAI
deployment. Run from the repository root: ``uv run python scripts/seed_workshop_demo.py``.
"""

import asyncio
import io
import logging
from datetime import UTC, datetime
from decimal import Decimal
from email.message import EmailMessage

import alloq_commons.entities  # noqa: F401  (registers the projects table)
import alloq_offer.workshop.worker  # noqa: F401  (registers the job handlers)
from alloq_offer.entities import IndustryEntity, OfferEntity, OfferStateEnum
from alloq_offer.workshop.configuration import WorkshopConfig
from alloq_offer.workshop.entities import SourceKind
from alloq_offer.workshop.jobs.runner import JobRunner
from alloq_offer.workshop.routes import WORKSHOP_ROUTE_PREFIX, WORKSHOP_ROUTE_SUFFIX
from alloq_offer.workshop.services.sources import add_source, request_initial_draft
from alloq_offer.workshop.services.workspace import open_workspace
from alloq_offer.workshop.storage import DbBlobStore
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import select

from appkit_commons.database.session import get_asyncdb_session
from appkit_commons.registry import service_registry

from app.configuration import configure

logger = logging.getLogger("seed_workshop_demo")

USER = "Demo"
TITLE = "Demo: Kundenportal Stadtwerke Musterstadt"
CUSTOMER = "Stadtwerke Musterstadt GmbH"
PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d49444154789c6360f8cfc000000301010018dd8db00000000049454e44ae426082"
)


def _docx() -> bytes:
    document = Document()
    document.add_heading("Lastenheft Kundenportal", level=1)
    for text in (
        "Heute melden viele Kunden ihre Zählerstände per Postkarte oder Telefon.",
        "Privatkunden sollen ihre Zählerstände online melden können.",
        "Privatkunden möchten ihre Abschläge selbst anpassen.",
        "Geschäftskunden wollen mehrere Verbrauchsstellen in einem Konto verwalten.",
        "Der Kundenservice muss Anliegen aus dem Portal ohne Medienbruch bearbeiten können.",
    ):
        document.add_paragraph(text)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Kanal"
    table.cell(0, 1).text = "Heute"
    table.cell(1, 0).text = "Zählerstand"
    table.cell(1, 1).text = "Postkarte, Telefon"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf() -> bytes:
    """A text PDF; the standard font has no umlauts, so they are spelled out."""
    lines = (
        "Fachkonzept Abrechnung",
        "Die Abrechnung muss Zaehlerstaende aus dem Portal automatisch plausibilisieren.",
        "Annahme: Das Abrechnungssystem bleibt SAP IS-U.",
    )
    writer = PdfWriter()
    font = writer._add_object(  # noqa: SLF001
        DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
    )
    page = writer.add_blank_page(595, 842)
    commands = "".join(
        f"BT /F1 11 Tf 40 {780 - 24 * i} Td ({line}) Tj ET\n"
        for i, line in enumerate(lines)
    )
    stream = DecodedStreamObject()
    stream.set_data(commands.encode())
    page[NameObject("/Contents")] = writer._add_object(stream)  # noqa: SLF001
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _pptx() -> bytes:
    deck = Presentation()
    for title, text, note in (
        (
            "Rechnungen",
            "Kunden wollen Rechnungen als PDF herunterladen.",
            "Wunsch aus dem Kickoff.",
        ),
        ("Kommunikation", "Kunden möchten Benachrichtigungen per E-Mail erhalten.", ""),
    ):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = title
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
        box.text_frame.text = text
        if note:
            slide.notes_slide.notes_text_frame.text = note
    deck.slides[0].shapes.add_picture(io.BytesIO(PNG), Inches(1), Inches(4))
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def _eml() -> bytes:
    message = EmailMessage()
    message["Subject"] = "Rückfragen der IT zum Kundenportal"
    message["From"] = "it@stadtwerke-musterstadt.example"
    message["To"] = "vertrieb@example.com"
    message["Message-ID"] = "<kundenportal-1@stadtwerke-musterstadt.example>"
    message.set_content(
        "Guten Tag,\n\n"
        "Ist eine Anmeldung über das bestehende Kundenkonto vorgesehen?\n"
        "Wir gehen davon aus, dass der Betrieb in der Azure Cloud erfolgt.\n"
        "Die IT-Sicherheit muss vor dem Go-live einen Penetrationstest durchführen.\n\n"
        "Viele Grüße\nIT Stadtwerke Musterstadt\n"
    )
    return message.as_bytes()


NOTES = (
    "Der Vertrieb möchte Tarifwechsel direkt im Portal anbieten. "
    "Brauchen Gewerbekunden eine Sammelrechnung?"
)
SOURCES = (
    ("Lastenheft Kundenportal.docx", SourceKind.DOCX, _docx),
    ("Fachkonzept Abrechnung.pdf", SourceKind.PDF, _pdf),
    ("Workshop Kickoff.pptx", SourceKind.PPTX, _pptx),
    ("Rückfragen IT.eml", SourceKind.EML, _eml),
)


async def _create_offer() -> int:
    async with get_asyncdb_session() as session:
        industry = await session.scalar(
            select(IndustryEntity.id).order_by(IndustryEntity.id)
        )
        if industry is None:
            raise SystemExit(
                "No industry found; create one on the industries page first."
            )
        offer = OfferEntity(
            title=TITLE,
            customer=CUSTOMER,
            person_days=Decimal(180),
            risk_buffer_days=Decimal(20),
            production_cost=Decimal(162000),
            recommended_price=Decimal(198000),
            key_accounter=USER,
            industry_id=industry,
            state=OfferStateEnum.IN_PROGRESS,
        )
        session.add(offer)
        await session.flush()
        await open_workspace(session, offer.id, USER)
        return offer.id


async def _run_jobs(runner: JobRunner) -> int:
    count = 0
    while await runner.run_once():
        count += 1
    return count


async def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    configure()
    config = service_registry().get(WorkshopConfig)
    if config.ai_provider != "demo" and not config.azure_openai.endpoint:
        raise SystemExit(
            "Set app.offer_workshop.ai_provider to demo or configure Azure OpenAI."
        )
    now = datetime.now(UTC)
    offer_id = await _create_offer()
    runner = JobRunner(config, offer_id=offer_id)
    store = DbBlobStore(config.max_blob_bytes)
    for filename, kind, build in SOURCES:
        await add_source(
            store,
            offer_id=offer_id,
            kind=kind,
            title=filename,
            filename=filename,
            mime="application/octet-stream",
            data=build(),
            user=USER,
            now=now,
            max_attempts=config.max_attempts,
        )
    await add_source(
        store,
        offer_id=offer_id,
        kind=SourceKind.TEXT,
        title="Notizen Telefonat Vertrieb",
        filename="",
        mime="text/plain",
        data=NOTES.encode(),
        user=USER,
        now=now,
        max_attempts=config.max_attempts,
    )
    logger.info("Read %d sources", await _run_jobs(runner))

    await request_initial_draft(
        get_asyncdb_session, offer_id, now=now, max_attempts=config.max_attempts
    )
    logger.info("Draft pipeline ran %d jobs", await _run_jobs(runner))

    logger.info(
        "Demo offer %s ready: %s%s%s",
        offer_id,
        WORKSHOP_ROUTE_PREFIX,
        offer_id,
        WORKSHOP_ROUTE_SUFFIX,
    )


if __name__ == "__main__":
    asyncio.run(main())
