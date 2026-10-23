# ruff: noqa: ANN001, ANN201, E501, PTH207
"""Build workshop/templates/quo_v1/layout.docx from the QUO customer template.

Keeps the cover and the static legal tail, removes the generated chapters and
replaces customer data by docxtpl placeholders. Run from the repository root.
"""

import copy
import glob
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TARGET = Path(
    "components/alloq-offer/src/alloq_offer/workshop/templates/quo_v1/layout.docx"
)
ANCHOR = "[[KAPITEL]]"
FIRST_GENERATED = "Management Summary"
FIRST_STATIC_TAIL = "Rechtliche Bestimmungen"

COVER_PLACEHOLDERS = {
    "MVP Referenzdatenbank": "{{ project_title }}",
    "TREYSTA Ingenieure Holding GmbH": "{{ customer_name }}",
    "Kurfürstendamm 29": "{{ customer_street }}",
    "10719 Berlin": "{{ customer_city }}",
    "- Im Folgenden „Treysta” oder „Auftraggeber“ genannt -": (
        "- Im Folgenden „{{ customer_short }}“ oder „Auftraggeber“ genannt -"
    ),
    "Leipzig, 19. Oktober 2026": "Leipzig, {{ offer_date }}",
    "Quote ID: QUO-33441-N7R0Z4": "Quote ID: {{ quote_id }}",
    "Der Auftragnehmer hält sich bis einschließlich zum 30.November.2026 an dieses Angebot gebunden.": (
        "Der Auftragnehmer hält sich bis einschließlich zum {{ validity_date }} an dieses Angebot gebunden."
    ),
}


def emit(*parts: object) -> None:
    sys.stdout.write(" ".join(str(part) for part in parts) + "\n")


def text_of(elem) -> str:
    return "".join(t.text or "" for t in elem.iter(qn("w:t")))


def set_text(paragraph, text: str) -> None:
    """Replace the paragraph content by one run that keeps the first run format."""
    first = paragraph.find(".//" + qn("w:r"))
    run = OxmlElement("w:r")
    if first is not None and first.find(qn("w:rPr")) is not None:
        run.append(copy.deepcopy(first.find(qn("w:rPr"))))
    node = OxmlElement("w:t")
    node.text = text
    node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    run.append(node)
    for child in list(paragraph):
        if child.tag != qn("w:pPr"):
            paragraph.remove(child)
    paragraph.append(run)


def heading1_index(kids, title: str) -> int:
    for i, kid in enumerate(kids):
        style = kid.find(qn("w:pPr") + "/" + qn("w:pStyle"))
        if (
            kid.tag == qn("w:p")
            and style is not None
            and style.get(qn("w:val")) == "berschrift1"
            and text_of(kid).strip() == title
        ):
            return i
    raise SystemExit(f"heading not found: {title}")


def toc_paragraph(doc: Document):
    title = doc.element.makeelement(qn("w:p"), {})
    ppr = OxmlElement("w:pPr")
    style = OxmlElement("w:pStyle")
    style.set(qn("w:val"), "berschriftohneNummerierung")
    ppr.append(style)
    title.append(ppr)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "Inhaltsverzeichnis"
    run.append(t)
    title.append(run)

    para = OxmlElement("w:p")
    for kind, payload in (
        ("begin", None),
        ("instr", ' TOC \\o "1-3" \\h \\z \\u '),
        ("separate", None),
        ("text", "Inhaltsverzeichnis wird beim Öffnen in Word aktualisiert."),
        ("end", None),
    ):
        r = OxmlElement("w:r")
        if kind == "instr":
            node = OxmlElement("w:instrText")
            node.text = payload
            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        elif kind == "text":
            node = OxmlElement("w:t")
            node.text = payload
        else:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        r.append(node)
        para.append(r)
    return title, para


def main() -> None:
    source = sys.argv[1] if len(sys.argv) > 1 else glob.glob("Templates/QUO*.docx")[0]
    doc = Document(source)
    body = doc.element.body
    kids = list(body.iterchildren())
    start = heading1_index(kids, FIRST_GENERATED)
    end = heading1_index(kids, FIRST_STATIC_TAIL)
    for kid in kids[start:end]:
        body.remove(kid)

    anchor = doc.element.makeelement(qn("w:p"), {})
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = ANCHOR
    run.append(node)
    anchor.append(run)
    toc_title, toc = toc_paragraph(doc)
    tail_first = kids[end]
    page_break = doc.element.makeelement(qn("w:p"), {})
    br_run = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    br_run.append(br)
    page_break.append(br_run)
    for elem in (toc_title, toc, page_break, anchor):
        tail_first.addprevious(elem)

    used: set[str] = set()
    for paragraph in body.iter(qn("w:p")):
        key = text_of(paragraph).strip()
        if key in COVER_PLACEHOLDERS:
            used.add(key)
            set_text(paragraph, COVER_PLACEHOLDERS[key])
    unmatched = set(COVER_PLACEHOLDERS) - used
    if unmatched:
        raise SystemExit(f"placeholders not matched: {sorted(unmatched)}")

    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)

    TARGET.parent.mkdir(parents=True, exist_ok=True)
    doc.save(TARGET)
    emit(f"written {TARGET}")


if __name__ == "__main__":
    main()
