"""Print an inventory of a DOCX template: paragraphs, styles, tables, fields."""

import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def emit(*parts: object) -> None:
    sys.stdout.write(" ".join(str(part) for part in parts) + "\n")


def main(path: str) -> None:
    doc = Document(path)
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[1]
        if tag == "p":
            style = child.find(qn("w:pPr") + "/" + qn("w:pStyle"))
            style_id = style.get(qn("w:val")) if style is not None else "-"
            text = "".join(t.text or "" for t in child.iter(qn("w:t")))
            fields = [i.text for i in child.iter(qn("w:instrText"))]
            sdt = len(list(child.iter(qn("w:sdt"))))
            emit(f"P [{style_id}] {text[:90]!r} fields={fields} sdt={sdt}")
        elif tag == "tbl":
            rows = child.findall(qn("w:tr"))
            first = [
                "".join(t.text or "" for t in c.iter(qn("w:t")))
                for c in rows[0].findall(qn("w:tc"))
            ]
            emit(f"TABLE rows={len(rows)} first_row={first}")
        elif tag == "sdt":
            emit("SDT block")
    for section in doc.sections:
        emit("HEADER", [p.text for p in section.header.paragraphs if p.text])
        emit("FOOTER", [p.text for p in section.footer.paragraphs if p.text])


if __name__ == "__main__":
    main(
        sys.argv[1]
        if len(sys.argv) > 1
        else str(Path("Templates").glob("QUO*.docx").__next__())
    )
